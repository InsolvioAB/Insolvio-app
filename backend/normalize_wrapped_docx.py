#!/usr/bin/env python3
"""Normalise a hard-wrapped plain-text style law DOCX (e.g. the Förmånsrättslag
file from the Teams share: title line, 'Departement:', 'Utfärdad:', 'Ändring
införd:', then lines wrapped at ~65 chars, blocks separated by empty paragraphs)
into the standard Regeringskansliet 'Fulltext' layout understood by
build_laws_from_docx.py.

  python3 normalize_wrapped_docx.py IN.docx OUT.docx

Rules (verified on formansrattslag.docx):
  * header = first non-empty lines up to the first blank line
  * block = consecutive non-empty paragraphs; blocks are separated by blanks
  * block whose first line matches '<n> [a-z] §' starts a section
  * block starting with '1. ', '- ' or 'a) ' is a list item of the current stycke
  * single-line block without final punctuation that is followed by a section
    start (or another heading) is a heading
  * everything after the last 'Övergångsbestämmelser' line is transitional text
"""
import re, sys
from docx import Document

NBSP = '\xa0'
SEC = re.compile(r'^(\d+)\s*([a-z]?)\s*§§?\*?\s*(.*)$', re.S)
LIST = re.compile(r'^(\d+\.|-|[a-z]\))\s')
SFS_ID = re.compile(r'^\d{4}:\d+$')


def blocks_of(lines):
    out, cur = [], []
    for t in lines:
        if t.strip():
            cur.append(t.strip())
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def join(lines):
    return re.sub(r'\s+', ' ', ' '.join(lines)).strip()


def main(src, dst):
    ps = [p.text for p in Document(src).paragraphs]
    # header
    hdr, i = [], 0
    while i < len(ps) and ps[i].strip():
        hdr.append(ps[i].strip().replace(NBSP, ' '))
        i += 1
    title = hdr[0]
    sfs = re.search(r'\((\d{4}:\d+)\)', title).group(1)
    meta = {'SFS nr': sfs}
    extra = []
    for h in hdr[1:]:
        if not re.match(r'^(Departement|Utfärdad|Ändring)', h):
            extra.append(h)
            continue
        k, _, v = h.partition(':')
        v = v.strip()
        k = k.strip()
        if k == 'Departement':
            meta['Departement/myndighet'] = v
        elif k.startswith('Utfärdad'):
            meta['Utfärdad'] = v
        elif k.startswith('Ändring'):
            meta['Ändrad'] = v
    # the first heading may be glued to the header (no blank line) -> handled below
    rest = ps[i:]
    ov = max(j for j, t in enumerate(rest) if t.strip() == 'Övergångsbestämmelser')
    body_blocks = ([extra] if extra else []) + blocks_of(rest[:ov])
    trans_blocks = blocks_of(rest[ov + 1:])

    doc = Document()
    doc.add_paragraph('SFS nr: {SFS nr}\nDepartement/myndighet: {Departement/myndighet}\nUtfärdad: {Utfärdad}\nÄndrad: {Ändrad}'.format(**meta))
    doc.add_paragraph('Innehåll:')
    doc.add_paragraph('Övergångsbestämmelser')

    def is_sec(b):
        return bool(SEC.match(b[0])) and re.match(r'^\d+\s*[a-z]?\s*§', b[0]) is not None

    def is_heading(idx):
        b = body_blocks[idx]
        if len(b) != 1 or is_sec(b) or LIST.match(b[0]) or b[0].startswith('/'):
            return False
        if b[0].rstrip().endswith(('.', ',', ';', ':', ')')):
            return False
        if idx + 1 >= len(body_blocks):
            return False
        nb = body_blocks[idx + 1]
        return is_sec(nb) or (len(nb) == 1 and not LIST.match(nb[0]) and not nb[0].endswith(('.', ',', ';', ':')) and idx + 2 < len(body_blocks) and is_sec(body_blocks[idx + 2]))

    cur_par = None  # list of text lines of the current stycke
    paras = []      # (kind, payload)
    for idx, b in enumerate(body_blocks):
        if is_sec(b):
            m = SEC.match(join(b))
            num, suf, txt = m.group(1), m.group(2), m.group(3)
            paras.append(('sec', f'{num}{(" " + suf) if suf else ""} §', txt))
        elif is_heading(idx):
            paras.append(('head', b[0]))
        elif LIST.match(b[0]) and paras and paras[-1][0] in ('sec', 'par'):
            # attach as list line to previous stycke
            last = paras[-1]
            if last[0] == 'sec':
                paras[-1] = ('sec', last[1], last[2] + '\n' + NBSP * 3 + join(b))
            else:
                paras[-1] = ('par', last[1] + '\n' + NBSP * 3 + join(b))
        else:
            paras.append(('par', join(b)))

    for p in paras:
        if p[0] == 'head':
            r = doc.add_paragraph().add_run(p[1])
            r.bold = True
        elif p[0] == 'sec':
            par = doc.add_paragraph()
            r = par.add_run(p[1] + NBSP * 3)
            r.bold = True
            par.add_run(p[2])
        else:
            doc.add_paragraph(p[1])

    doc.add_paragraph('Övergångsbestämmelser')
    out = []
    for b in trans_blocks:
        t = join(b)
        if SFS_ID.match(t):
            out.append(t)
        elif out and not SFS_ID.match(out[-1]) and not re.match(r'^\d+\.\s', t) and out[-1].rstrip().endswith((',', ':', 'upphäves')):
            out[-1] += '\n' + NBSP * 3 + t
        elif out and not SFS_ID.match(out[-1]) and not re.match(r'^\d+\.\s', t) and '\n' in out[-1] and out[-1].rstrip().endswith((',',)):
            out[-1] += '\n' + NBSP * 3 + t
        else:
            out.append(t)
    for t in out:
        doc.add_paragraph(t)
    doc.save(dst)
    print('sections:', sum(1 for p in paras if p[0] == 'sec'), 'headings:', [p[1] for p in paras if p[0] == 'head'], 'trans paras:', len(out))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
